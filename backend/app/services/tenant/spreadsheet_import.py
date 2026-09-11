"""Read a CSV or XLSX file into the shape a spreadsheet document holds.

The inverse of :mod:`app.services.export.spreadsheet`, and deliberately its
neighbour: the two have to agree about how a fill, a border or a currency
format is spelled, and they only stay in agreement while they are read
together. Every conversion here names the export function it reverses.

Parsing happens on the server for the same reason rendering does — the
workbook libraries live here, and the result goes through
``normalize_spreadsheet_content``, so an imported sheet arrives in exactly the
shape a created one does rather than in a second shape that has to be kept
in step with the first.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

from app.core.messages import DocumentMessages
from app.services.tenant.documents_spreadsheet import (
    MAX_COLS,
    MAX_ROWS,
    MAX_SHEETS,
    DocumentContentError,
    normalize_spreadsheet_content,
)

# The same conversions the renderer uses, read the other way.
_PX_TO_POINTS = 0.75
_PX_PER_WIDTH_UNIT = 7.0
_VALIGN_FROM_XLSX = {"top": "top", "center": "middle", "bottom": "bottom"}
_ALIGN_VALUES = {"left", "center", "right"}

# The size Excel gives a cell nobody styled.
_XLSX_DEFAULT_FONT_PT = 11.0

# The grid a new sheet comes with. An imported sheet is at least this big, so
# a three-row CSV opens as something there is room to work in rather than as a
# three-row grid.
_FLOOR_ROWS = 100
_FLOOR_COLS = 26

# The point past which a file is not a spreadsheet somebody is editing by
# hand. Reached, the import is refused rather than trimmed: a workbook that
# came back missing the rows past some line, reported as imported, is worse
# than one that did not come back at all.
MAX_IMPORT_CELLS: int = 500_000

# The most coordinates one import will look at. A sheet declares the rectangle
# it covers, and that rectangle can be enormously larger than the cells it
# actually holds — every coordinate inside one still costs a read. Measured
# from the declared rectangle before anything is read, so a sheet that would
# take too long never starts.
MAX_IMPORT_SCAN: int = 2_000_000


def parse_spreadsheet_file(filename: str, data: bytes) -> list[dict[str, Any]]:
    """The sheets a file holds, in canonical workbook form.

    Raises :class:`DocumentContentError` for anything that is not a
    spreadsheet this can read.
    """
    name = (filename or "").strip()
    lower = name.lower()
    base = name.rsplit("/", 1)[-1].rsplit(".", 1)[0] or "Imported"

    if lower.endswith(".csv") or lower.endswith(".tsv"):
        raw = [_parse_csv(data, base, tab=lower.endswith(".tsv"))]
    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        raw = _parse_xlsx(data)
    else:
        raise DocumentContentError(DocumentMessages.SPREADSHEET_UNREADABLE_FILE)

    if not raw:
        raise DocumentContentError(DocumentMessages.SPREADSHEET_UNREADABLE_FILE)

    # One trip through the normalizer the create/patch paths use, so an
    # imported sheet is the same kind of object as any other.
    normalized = normalize_spreadsheet_content(
        {"schema_version": 3, "kind": "spreadsheet", "sheets": raw}
    )
    return normalized["sheets"]


# ── CSV ──────────────────────────────────────────────────────────────────────


def _decode(data: bytes) -> str:
    """Text from bytes, preferring UTF-8 and never failing on a file that
    is not. A spreadsheet exported by an older tool is often cp1252, and a
    mangled accent is a better outcome than a refused import."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(data: bytes, name: str, *, tab: bool) -> dict[str, Any]:
    text = _decode(data)
    delimiter = "\t" if tab else ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    cells: dict[str, Any] = {}
    rows = 0
    cols = 0
    for r, record in enumerate(reader):
        if r >= MAX_ROWS:
            raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
        rows = r + 1
        for c, value in enumerate(record):
            if c >= MAX_COLS:
                raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
            cols = max(cols, c + 1)
            if value == "":
                continue
            if len(cells) >= MAX_IMPORT_CELLS:
                raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
            cells[f"{r}:{c}"] = _scalar(value)
    return {
        "name": name,
        "cells": cells,
        "dimensions": _dimensions(rows, cols),
    }


# What a number looks like, matching ``coerceScalar`` in
# ``frontend/src/lib/spreadsheet/csv.ts``. A field pasted from the clipboard
# and the same field read from a file have to become the same value.
_NUMERIC_RE = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


def _scalar(text: str) -> Any:
    """A CSV field as the value it spells.

    A leading ``=`` stays a string — that is how this app stores a formula,
    and the grid evaluates it on read. A leading zero also keeps the field as
    text: ``00123`` is a part number, a postcode or an extension far more often
    than it is the number 123, and turning it into one cannot be undone.
    """
    trimmed = text.strip()
    if trimmed == "":
        return ""
    if trimmed.startswith("="):
        return trimmed
    lower = trimmed.lower()
    if lower == "true":
        return True
    if lower == "false":
        return False
    if _NUMERIC_RE.match(trimmed):
        if trimmed.startswith("0") and not trimmed.startswith("0.") and trimmed != "0":
            return trimmed
        try:
            number = float(trimmed)
        except ValueError:
            return text
        if number == int(number) and "." not in trimmed and "e" not in lower:
            return int(number)
        return number
    return text


# ── XLSX ─────────────────────────────────────────────────────────────────────


def _parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    try:
        # ``data_only=False`` keeps a formula as its ``=`` text, matching how
        # the renderer writes one and how the grid stores it.
        workbook = load_workbook(
            io.BytesIO(data), data_only=False, keep_links=False, rich_text=False
        )
    except Exception as exc:  # openpyxl raises a zoo of types on bad input
        raise DocumentContentError(
            DocumentMessages.SPREADSHEET_UNREADABLE_FILE
        ) from exc

    if len(workbook.worksheets) > MAX_SHEETS:
        raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
    _refuse_unreadable_shape(workbook.worksheets)

    budget = MAX_IMPORT_CELLS
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        sheet, used = _parse_worksheet(worksheet, budget)
        budget -= used
        sheets.append(sheet)
    return sheets


def _refuse_unreadable_shape(worksheets: list[Worksheet]) -> None:
    """Refuse a workbook whose declared shape is more than one import reads.

    Checked across every sheet before any is read, so the cost of deciding is
    the same whatever the file claims to hold.
    """
    scan = 0
    for ws in worksheets:
        rows = ws.max_row or 0
        cols = ws.max_column or 0
        if rows > MAX_ROWS or cols > MAX_COLS:
            raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
        scan += rows * cols
        if scan > MAX_IMPORT_SCAN:
            raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)


def _parse_worksheet(ws: Worksheet, budget: int) -> tuple[dict[str, Any], int]:
    cells: dict[str, Any] = {}
    cell_styles: dict[str, Any] = {}
    rows = 0
    cols = 0
    used = 0

    for row in ws.iter_rows():
        for cell in row:
            value = _cell_value(cell)
            style = _style_of(cell)
            if value is None and style is None:
                continue
            r = cell.row - 1
            c = cell.column - 1
            used += 1
            if used > budget:
                # Better to say a file is too big than to hand back some of it
                # and call that the file.
                raise DocumentContentError(DocumentMessages.SPREADSHEET_FILE_TOO_LARGE)
            rows = max(rows, r + 1)
            cols = max(cols, c + 1)
            if value is not None:
                cells[f"{r}:{c}"] = value
            if style is not None:
                cell_styles[f"{r}:{c}"] = style

    return _sheet(ws, cells, cell_styles, rows, cols), used


def _sheet(
    ws: Worksheet,
    cells: dict[str, Any],
    cell_styles: dict[str, Any],
    rows: int,
    cols: int,
) -> dict[str, Any]:
    frozen = _frozen(ws)
    return {
        "name": ws.title,
        "cells": cells,
        "cellStyles": cell_styles,
        "columns": _columns(ws),
        "rows": _rows(ws),
        "frozen": frozen,
        "dimensions": _dimensions(
            max(rows, ws.max_row or 0, frozen["rows"] + 1),
            max(cols, ws.max_column or 0, frozen["cols"] + 1),
        ),
        **({"hidden": True} if ws.sheet_state != "visible" else {}),
    }


def _dimensions(rows: int, cols: int) -> dict[str, int]:
    """A grid big enough for the content and big enough to work in."""
    return {
        "rows": min(max(rows, _FLOOR_ROWS), MAX_ROWS),
        "cols": min(max(cols, _FLOOR_COLS), MAX_COLS),
    }


def _cell_value(cell: Cell) -> Any:
    """The scalar a cell holds, or ``None`` for an empty one.

    Dates arrive as ``datetime`` and are stored as ISO text — the grid's
    values are text, numbers and booleans, and the cell's number format is
    what makes one of them read as a date.
    """
    value = cell.value
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time() == time(0, 0)
            else value.isoformat(sep=" ")
        )
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


# ── styles: the inverse of ``_apply_style`` ──────────────────────────────────


def _style_of(cell: Cell) -> dict[str, Any] | None:
    """A cell's entry in ``cellStyles``: its look under ``style``, its number
    format under ``format`` — the wrapper every layer of formatting uses.
    ``None`` when the cell wears the default, so an untouched one costs
    nothing to import."""
    style: dict[str, Any] = {}

    font = cell.font
    if font is not None:
        if font.bold:
            style["bold"] = True
        if font.italic:
            style["italic"] = True
        if font.underline:
            style["underline"] = True
        if font.strike:
            style["strike"] = True
        color = _hex(getattr(font.color, "rgb", None))
        # Excel writes the default body colour explicitly; carrying it over
        # would pin every cell to black and defeat the theme.
        if color and color not in ("#000000",):
            style["color"] = color
        # 11pt is the size a cell has when nobody chose one — Excel writes it
        # on every cell, so carrying it over would set an explicit size on a
        # whole imported sheet. A cell genuinely set to 11pt reads as default,
        # which is the lesser of the two wrongs.
        if (
            isinstance(font.size, (int, float))
            and font.size
            and abs(font.size - _XLSX_DEFAULT_FONT_PT) > 0.01
        ):
            px = round(font.size / _PX_TO_POINTS)
            if px:
                style["fontSize"] = px

    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        rgb = _hex(getattr(fill.start_color, "rgb", None))
        # A white fill is Excel's way of saying "no fill" as often as it is a
        # deliberate one; treating it as deliberate paints over the theme.
        if rgb and rgb != "#ffffff":
            style["fill"] = rgb

    alignment = cell.alignment
    if alignment is not None:
        if alignment.horizontal in _ALIGN_VALUES:
            style["align"] = alignment.horizontal
        valign = _VALIGN_FROM_XLSX.get(alignment.vertical or "")
        if valign:
            style["valign"] = valign

    border = _border_of(cell)
    if border:
        style["border"] = border

    entry: dict[str, Any] = {}
    if style:
        entry["style"] = style
    fmt = _format_of(cell.number_format)
    if fmt:
        entry["format"] = fmt
    return entry or None


def _border_of(cell: Cell) -> dict[str, Any]:
    out: dict[str, Any] = {}
    border = cell.border
    if border is None:
        return out
    for edge in ("top", "right", "bottom", "left"):
        side = getattr(border, edge, None)
        if side is None or not side.style:
            continue
        spec: dict[str, Any] = {"style": side.style}
        color = _hex(getattr(side.color, "rgb", None))
        if color:
            spec["color"] = color
        out[edge] = spec
    return out


def _hex(value: Any) -> str | None:
    """openpyxl's ``AARRGGBB`` -> ``#rrggbb``. The inverse of ``_argb``.

    A theme or indexed colour has no literal rgb and is left to the theme.
    """
    if not isinstance(value, str) or len(value) != 8:
        return None
    body = value[2:]
    if any(ch not in "0123456789abcdefABCDEF" for ch in body):
        return None
    return f"#{body.lower()}"


def _format_of(number_format: Any) -> dict[str, Any] | None:
    """A number format string as the grid's format object — the inverse of
    ``_number_format``. Anything this does not recognize is left plain, which
    is what the grid does with a format it cannot name."""
    if not isinstance(number_format, str):
        return None
    text = number_format.strip()
    if not text or text.lower() == "general":
        return None

    head = text.split(";", 1)[0]
    # Only the run of digit placeholders straight after the point is the
    # decimal count; a currency's trailing "USD" is not.
    fraction = re.match(r"[0#]*", head.split(".", 1)[1]) if "." in head else None
    decimals = min(len(fraction.group(0)) if fraction else 0, 10)
    grouping = "#,##" in head

    if head.endswith("%"):
        return {"type": "percent", "decimals": decimals}
    if any(token in head for token in ("yy", "mmm", "dd/", "-mm-", "/dd")):
        return {"type": "date", "pattern": "iso"}
    if '"' in head or "$" in head or "€" in head or "£" in head:
        currency = _currency_of(head)
        out: dict[str, Any] = {
            "type": "currency",
            "decimals": decimals,
            "grouping": grouping,
        }
        if currency:
            out["currency"] = currency
        if "[Red]" in text:
            out["negatives"] = "red"
        return out
    if "0" in head or "#" in head:
        return {"type": "fixed", "decimals": decimals, "grouping": grouping}
    return None


def _currency_of(head: str) -> str | None:
    """The code a currency format names, when it names one literally."""
    if '"' in head:
        parts = head.split('"')
        if len(parts) >= 2 and parts[1].strip():
            return parts[1].strip()[:8]
    for symbol, code in (("$", "USD"), ("€", "EUR"), ("£", "GBP")):
        if symbol in head:
            return code
    return None


# ── sheet-level geometry ─────────────────────────────────────────────────────


def _columns(ws: Worksheet) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for letter, dim in (ws.column_dimensions or {}).items():
        try:
            index = column_index_from_string(letter) - 1
        except Exception:
            continue
        if index < 0 or index >= MAX_COLS:
            continue
        entry: dict[str, Any] = {}
        if isinstance(dim.width, (int, float)) and dim.width:
            entry["width"] = round(dim.width * _PX_PER_WIDTH_UNIT)
        if dim.hidden:
            entry["hidden"] = True
        if entry:
            out[str(index)] = entry
    return out


def _rows(ws: Worksheet) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for number, dim in (ws.row_dimensions or {}).items():
        index = int(number) - 1
        if index < 0 or index >= MAX_ROWS:
            continue
        entry: dict[str, Any] = {}
        if isinstance(dim.height, (int, float)) and dim.height:
            entry["height"] = round(dim.height / _PX_TO_POINTS)
        if dim.hidden:
            entry["hidden"] = True
        if entry:
            out[str(index)] = entry
    return out


def _frozen(ws: Worksheet) -> dict[str, int]:
    """Frozen panes from the top-left unfrozen cell openpyxl records."""
    pane = ws.freeze_panes
    if not isinstance(pane, str) or not pane:
        return {"rows": 0, "cols": 0}
    try:
        letter, row = coordinate_from_string(pane)
        return {
            "rows": max(row - 1, 0),
            "cols": max(column_index_from_string(letter) - 1, 0),
        }
    except Exception:
        return {"rows": 0, "cols": 0}
